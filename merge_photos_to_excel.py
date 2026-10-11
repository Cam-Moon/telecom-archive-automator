import io
import os
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Font
from PIL import Image as PILImage

EXCEL_FILE = "Vonage_SMS_History_Full_Tabs.xlsx"
PHOTOS_DIR = "Vonage_Photos"
OUTPUT_FILE = "Vonage_SMS_History_With_Photos.xlsx"

print("Scanning local photos...")
photo_files = {}
for fname in os.listdir(PHOTOS_DIR):
  if not fname.startswith(".") and os.path.isfile(
      os.path.join(PHOTOS_DIR, fname)
  ):
    parts = fname.split("_")
    if len(parts) >= 2:
      key = f"{parts[0]}_{parts[1]}"
      if key not in photo_files:
        photo_files[key] = []
      photo_files[key].append(fname)

print(
    f"Loaded {sum(len(v) for v in photo_files.values())} photos. Opening"
    " Excel..."
)
wb = openpyxl.load_workbook(EXCEL_FILE)

for sheet_name in wb.sheetnames:
  ws = wb[sheet_name]
  if ws.max_row < 2:
    continue

  # Column E header (Thumbnail)
  ws.cell(row=1, column=5, value="Thumbnail")
  ws.cell(row=1, column=5).font = Font(bold=True)
  ws.column_dimensions["E"].width = 16

  # Column F header (Link)
  ws.cell(row=1, column=6, value="Image Link")
  ws.cell(row=1, column=6).font = Font(bold=True)
  ws.column_dimensions["F"].width = 35

  for row in range(2, ws.max_row + 1):
    msg_val = str(ws.cell(row=row, column=4).value or "")
    date_val = str(ws.cell(row=row, column=1).value or "").strip()

    if "[Image/Attachment]" in msg_val:
      key = f"{sheet_name}_{date_val}"

      if key in photo_files and len(photo_files[key]) > 0:
        matched_filename = photo_files[key].pop(0)
        full_image_path = os.path.join(PHOTOS_DIR, matched_filename)

        try:
          # Build a lightweight thumbnail in memory
          with PILImage.open(full_image_path) as pil_img:
            pil_img.thumbnail((75, 75))
            img_bytes = io.BytesIO()
            pil_img.convert("RGB").save(img_bytes, format="JPEG", quality=70)
            img_bytes.seek(0)

          # Place thumbnail in Column E
          xl_img = OpenpyxlImage(img_bytes)
          xl_img.width = 55
          xl_img.height = 55
          ws.add_image(xl_img, f"E{row}")
          ws.row_dimensions[row].height = 48

          # Add relative clickable link in Column F
          link_cell = ws.cell(row=row, column=6)
          link_cell.value = f"Open Full Image ({matched_filename})"
          link_cell.hyperlink = f"{PHOTOS_DIR}/{matched_filename}"
          link_cell.font = Font(color="0000FF", underline="single")
        except Exception as e:
          print(f"Skipping {matched_filename}: {e}")

print("Saving merged workbook...")
wb.save(OUTPUT_FILE)
print(f"Complete! Saved as '{OUTPUT_FILE}'.")